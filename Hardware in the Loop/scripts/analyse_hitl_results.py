#!/usr/bin/env python3
"""
HITL R54 Results Analyser — generates Excel workbook + publication-quality PNGs.

Usage:
  python3 "Hardware in the Loop/scripts/analyse_hitl_results.py"
  python3 "Hardware in the Loop/scripts/analyse_hitl_results.py" --summary path/to/file.json
"""

import argparse
import json
import math
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import matplotlib.ticker as mticker
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Publication style ─────────────────────────────────────────────────────────
plt.rcParams.update({
    "figure.facecolor":  "white",
    "axes.facecolor":    "#F7F9FC",
    "axes.edgecolor":    "#CCCCCC",
    "axes.linewidth":    1.3,
    "axes.labelcolor":   "#1E293B",
    "axes.labelsize":    13,
    "axes.titlesize":    14,
    "axes.titleweight":  "bold",
    "axes.titlepad":     10,
    "xtick.color":       "#475569",
    "ytick.color":       "#475569",
    "xtick.labelsize":   11,
    "ytick.labelsize":   11,
    "grid.color":        "#E2E8F0",
    "grid.linewidth":    0.9,
    "grid.linestyle":    "--",
    "text.color":        "#1E293B",
    "lines.linewidth":   2.2,
    "legend.framealpha": 0.92,
    "legend.edgecolor":  "#CBD5E1",
    "legend.fontsize":   10,
    "savefig.dpi":       180,
    "savefig.facecolor": "white",
    "savefig.bbox":      "tight",
    "font.family":       "DejaVu Sans",
})

# ── Colours ───────────────────────────────────────────────────────────────────
C_SUCCESS  = "#2563EB"
C_TIMEOUT  = "#D97706"
C_CRASH    = "#DC2626"
C_STUCK    = "#7C3AED"
C_RETURN   = "#60A5FA"
C_GREEN    = "#16A34A"

OUTCOME_COL = {
    "SUCCESS":   C_SUCCESS,
    "TIMEOUT":   C_TIMEOUT,
    "CRASH_ALT": C_CRASH,
    "CRASH_LAT": C_CRASH,
    "STUCK":     C_STUCK,
}

RESULTS_DIR = Path(__file__).resolve().parents[1] / "results"


# ── Helpers ───────────────────────────────────────────────────────────────────
def find_latest_summary() -> Path:
    candidates = sorted(
        RESULTS_DIR.glob("hitl_r54_*_summary_*.json"),
        key=lambda p: p.stat().st_mtime)
    if not candidates:
        sys.exit(f"[ERROR] No summary JSON found in {RESULTS_DIR}")
    return candidates[-1]


def load_csv(path: str) -> "pd.DataFrame | None":
    p = Path(path)
    if not p.exists():
        return None
    try:
        df = pd.read_csv(p)
    except Exception:
        return None
    # Replace MAVROS-EKF speed artifacts with position-derived speed.
    # When the EKF reports a single constant impossible value (>4.5 m/s on every
    # step) the column is garbage; recompute from consecutive position deltas.
    spd = df["speed_ms"].values
    if len(spd) > 1 and np.all(spd == spd[0]) and spd[0] > 4.5:
        dt = np.diff(df["t_s"].values, prepend=df["t_s"].iloc[0])
        dt = np.where(dt <= 0, 0.1, dt)
        dx = np.diff(df["x_m"].values, prepend=df["x_m"].iloc[0])
        dy = np.diff(df["y_m"].values, prepend=df["y_m"].iloc[0])
        dz = np.diff(df["z_m"].values, prepend=df["z_m"].iloc[0])
        df["speed_ms"] = np.sqrt(dx**2 + dy**2 + dz**2) / dt
        df["speed_ms"] = df["speed_ms"].clip(upper=6.0)
    return df


def col(ep: dict) -> str:
    return OUTCOME_COL.get(ep["result"], "#64748B")


def annotate_bars(ax, bars, fmt="{:.1f}", dy_frac=0.02):
    ylim = ax.get_ylim()
    dy = (ylim[1] - ylim[0]) * dy_frac
    for bar in bars:
        h = bar.get_height()
        if h > 0:
            ax.text(bar.get_x() + bar.get_width() / 2, h + dy,
                    fmt.format(h), ha="center", va="bottom",
                    fontsize=9.5, fontweight="bold", color="#1E293B")


def outcome_legend(ax, episodes):
    seen = {}
    for ep in episodes:
        r = ep["result"]
        if r not in seen:
            seen[r] = mpatches.Patch(color=col(ep), label=r)
    ax.legend(handles=list(seen.values()), fontsize=9, framealpha=0.9)


# ── Figure 1 — Mission progress (per-episode tunnel-position vs time) ─────────
def fig_mission_progress(eps_done: list, meta: dict, out: Path):
    n    = len(eps_done)
    ncol = min(3, n)
    nrow = math.ceil(n / ncol)
    fig, axes = plt.subplots(nrow, ncol,
                             figsize=(ncol * 5.8, nrow * 4.2), squeeze=False)
    n_ok = sum(1 for e in meta["episodes"] if e["result"] == "SUCCESS")
    n_ep = meta["n_episodes"]

    fig.suptitle("R54 HITL — Mission Progress · Round-Trip Tunnel Navigation",
                 fontsize=16, fontweight="bold")
    fig.text(0.5, 0.975,
             f"Domain Randomisation — R54 Policy  ·  RDDRONE-FMUK66 Hardware FC  ·  "
             f"{n_ok}/{n_ep} successful  ·  Corridor ±2.5 m  ·  Turn point 95 m",
             ha="center", fontsize=10.5, color="#475569")

    for i, ep in enumerate(eps_done):
        r, c = divmod(i, ncol)
        ax   = axes[r][c]
        df   = ep["df"]
        t    = df["t_s"].values
        x    = df["x_m"].values          # raw x (goes negative on return)
        pos  = np.abs(x)                 # "distance into tunnel" for display

        ob   = df[df["phase"] == "OUTBOUND"]
        rt   = df[df["phase"] == "RETURN"]

        ax.fill_between(t, pos, alpha=0.13, color=C_SUCCESS, zorder=1)
        ax.plot(t, pos, color=C_SUCCESS, lw=2.4, zorder=2, label="Tunnel position (m)")
        ax.axhline(95, color="#64748B", lw=1.3, ls="--", label="Turn-point (95 m)", zorder=1)

        if not rt.empty:
            t_ret = rt["t_s"].iloc[0]
            ax.axvline(t_ret, color=C_RETURN, lw=1.1, ls=":", alpha=0.8, zorder=3)
            ax.text(t_ret + 1.5, 4, "← Return", color=C_RETURN,
                    fontsize=9, style="italic")
        if not ob.empty and not rt.empty:
            t_ob  = ob["t_s"].iloc[-1]
            x_ob  = pos[df["phase"].values == "OUTBOUND"][-1]
            ax.annotate("Outbound\ncomplete",
                        xy=(t_ob, x_ob),
                        xytext=(max(0, t_ob - 14), x_ob - 18),
                        fontsize=8.5, color=C_SUCCESS,
                        arrowprops=dict(arrowstyle="->", color=C_SUCCESS, lw=1.2))

        status = "✓ SUCCESS" if ep["result"] == "SUCCESS" else f"✗ {ep['result']}"
        ax.set_title(f"(A{ep['attempt']:02d})  {status}  ({ep['duration_s']:.1f} s)",
                     color=col(ep), fontsize=12, fontweight="bold", pad=6)
        ax.set_xlabel("Time (s)")
        ax.set_ylabel("Tunnel position (m)")
        ax.set_ylim(bottom=-2)
        ax.legend(loc="upper left", fontsize=8.5)
        ax.grid(True, zorder=0)

    for i in range(n, nrow * ncol):
        r, c = divmod(i, ncol)
        axes[r][c].set_visible(False)

    plt.tight_layout(rect=[0, 0, 1, 0.96])
    plt.savefig(out)
    plt.close(fig)
    print(f"  Saved: {out.name}")


# ── Figure 2 — Episode metrics comparison bar chart ───────────────────────────
def fig_metrics(meta: dict, out: Path):
    eps  = [e for e in meta["episodes"] if e["result"] != "pending"]
    n    = len(eps)
    lbls = [f"A{e['attempt']:02d}" for e in eps]
    cols = [col(e) for e in eps]
    x    = np.arange(n)

    # Recompute avg speed from CSV where available (summary may have anomalies)
    speeds_avg  = []
    speeds_peak = []
    for ep in eps:
        df = ep.get("df")
        if df is not None and not df.empty:
            spd = df["speed_ms"].values
            speeds_avg.append(float(np.mean(spd)))
            speeds_peak.append(float(np.max(spd)))
        else:
            speeds_avg.append(ep["avg_speed_ms"])
            speeds_peak.append(ep["peak_speed_ms"])

    panels = [
        ("Duration\n(s)",          [e["duration_s"]  for e in eps],    "{:.1f}"),
        ("Avg Speed\n(m/s)",        speeds_avg,                          "{:.3f}"),
        ("Peak Speed\n(m/s)",       speeds_peak,                         "{:.3f}"),
        ("Min Obstacle\nDist. (m)", [e["min_depth_m"] for e in eps],    "{:.3f}"),
        ("Safety Events\n(col+nm)", [e["collisions"] + e["near_misses"] for e in eps], "{:.0f}"),
    ]

    fig, axes = plt.subplots(1, len(panels),
                             figsize=(len(panels) * 3.2 + 0.5, 5.8))
    n_ok = sum(1 for e in eps if e["result"] == "SUCCESS")
    fig.suptitle("R54 HITL — Episode Metrics Comparison",
                 fontsize=16, fontweight="bold")
    fig.text(0.5, 0.93,
             f"RDDRONE-FMUK66 · PX4 HITL Mode  ·  {n_ok}/{n} successful  ·  Colour = outcome",
             ha="center", fontsize=10.5, color="#475569")

    for ax, (title, vals, fmt) in zip(axes, panels):
        bars = ax.bar(x, vals, color=cols, edgecolor="white", linewidth=0.7,
                      width=0.68, zorder=2)
        ax.set_xticks(x)
        ax.set_xticklabels(lbls, rotation=45 if n > 6 else 0,
                           ha="right" if n > 6 else "center", fontsize=9.5)
        ax.set_title(title, fontsize=12, fontweight="bold")
        mx = max(vals) if vals else 1
        ax.set_ylim(0, mx * 1.28 if mx > 0 else 1)
        annotate_bars(ax, bars, fmt=fmt, dy_frac=0.025)
        ax.grid(True, axis="y", zorder=0)
        ax.set_axisbelow(True)

    handles = [mpatches.Patch(color=v, label=k) for k, v in OUTCOME_COL.items()
               if any(e["result"] == k for e in eps)]
    fig.legend(handles=handles, loc="lower center", ncol=len(handles),
               framealpha=0.92, fontsize=9.5, bbox_to_anchor=(0.5, -0.04))

    plt.tight_layout(rect=[0, 0.04, 1, 0.91])
    plt.savefig(out)
    plt.close(fig)
    print(f"  Saved: {out.name}")


# ── Figure 3 — Top-down trajectory overlay ────────────────────────────────────
def fig_trajectories(eps_done: list, meta: dict, out: Path):
    fig, ax = plt.subplots(figsize=(15, 5.0))
    n_ok = sum(1 for e in meta["episodes"] if e["result"] == "SUCCESS")
    n_ep = meta["n_episodes"]
    fig.suptitle("R54 HITL — Top-down Trajectory Overlay  ·  All Episodes",
                 fontsize=16, fontweight="bold")
    fig.text(0.5, 0.94,
             f"RDDRONE-FMUK66 · PX4 HITL Mode  ·  {n_ok}/{n_ep} successful  ·  "
             "Corridor ±2.5 m  ·  Outbound solid, Return dashed",
             ha="center", fontsize=10.5, color="#475569")

    ax.axhspan(-2.5, 2.5, color="#EEF2FF", zorder=0, label="Corridor ±2.5 m")
    ax.axhline( 2.5, color="#94A3B8", lw=1.3, ls="--", alpha=0.8)
    ax.axhline(-2.5, color="#94A3B8", lw=1.3, ls="--", alpha=0.8)
    ax.axhline(0,    color="#CBD5E1", lw=0.8, zorder=0)
    ax.axvline(95,   color="#94A3B8", lw=1.0, ls=":", alpha=0.7, zorder=1)
    ax.text(96, 2.3, "Turn\n95 m", fontsize=8.5, color="#64748B", va="top")

    legend_handles = []
    for ep in eps_done:
        df  = ep["df"]
        c   = col(ep)
        lbl = f"A{ep['attempt']:02d} — {ep['result']}"
        ob  = df[df["phase"] == "OUTBOUND"]
        rt  = df[df["phase"] == "RETURN"]
        if not ob.empty:
            ax.plot(ob["x_m"], ob["y_m"], color=c, lw=2.0, alpha=0.85, zorder=2)
        if not rt.empty:
            ax.plot(rt["x_m"], rt["y_m"], color=c, lw=1.5, ls="--",
                    alpha=0.70, zorder=2)
        # markers
        ax.scatter([df["x_m"].iloc[0]],  [df["y_m"].iloc[0]],
                   color=c, s=25, zorder=4)
        ax.scatter([df["x_m"].iloc[-1]], [df["y_m"].iloc[-1]],
                   color=c, s=65, marker="*", zorder=5)
        ax.text(df["x_m"].iloc[-1] + 0.8, df["y_m"].iloc[-1],
                f"A{ep['attempt']:02d}", fontsize=7.5, color=c, va="center")
        legend_handles.append(mpatches.Patch(color=c, label=lbl))

    ax.set_xlim(-3, 107)
    ax.set_ylim(-3.6, 3.6)
    ax.set_xlabel("Tunnel x-position (m)", fontsize=13)
    ax.set_ylabel("Lateral offset y (m)", fontsize=13)
    ax.legend(handles=legend_handles, loc="upper left", fontsize=8.5,
              framealpha=0.92, ncol=max(1, len(eps_done) // 5))
    ax.grid(True, zorder=0)
    plt.tight_layout(rect=[0, 0, 1, 0.92])
    plt.savefig(out)
    plt.close(fig)
    print(f"  Saved: {out.name}")


# ── Figure 4 — Speed & altitude profiles ──────────────────────────────────────
def fig_speed_altitude(eps_done: list, meta: dict, out: Path):
    n    = len(eps_done)
    ncol = min(3, n)
    nrow = math.ceil(n / ncol)
    fig, axes = plt.subplots(nrow * 2, ncol,
                             figsize=(ncol * 5.5, nrow * 5.5 + 1.0),
                             squeeze=False)
    n_ok = sum(1 for e in meta["episodes"] if e["result"] == "SUCCESS")
    n_ep = meta["n_episodes"]
    fig.suptitle("R54 HITL — Speed & Altitude Profiles",
                 fontsize=16, fontweight="bold")
    fig.text(0.5, 0.985,
             f"RDDRONE-FMUK66 · PX4 HITL Mode  ·  {n_ok}/{n_ep} successful",
             ha="center", fontsize=10.5, color="#475569")

    for i, ep in enumerate(eps_done):
        col_i = i % ncol
        row_s = (i // ncol) * 2
        ax_s  = axes[row_s    ][col_i]
        ax_a  = axes[row_s + 1][col_i]
        df    = ep["df"]
        t     = df["t_s"].values

        # Return phase shading
        rt_mask = df["phase"].values == "RETURN"
        if rt_mask.any():
            t_ret = df.loc[rt_mask, "t_s"].iloc[0]
            for ax in (ax_s, ax_a):
                ax.axvspan(t_ret, t[-1], color="#EFF6FF", alpha=0.5, zorder=0)

        # Speed
        ax_s.plot(t, df["speed_ms"], color=C_SUCCESS, lw=2.2, zorder=2)
        ax_s.fill_between(t, df["speed_ms"], alpha=0.13, color=C_SUCCESS, zorder=1)
        ax_s.set_ylabel("Speed (m/s)", fontsize=10)
        ax_s.set_ylim(bottom=0)
        ax_s.grid(True, zorder=0)

        # Altitude
        ax_a.plot(t, df["z_m"], color=C_GREEN, lw=2.2, zorder=2)
        ax_a.fill_between(t, df["z_m"], alpha=0.10, color=C_GREEN, zorder=1)
        ax_a.axhline(1.5, color="#94A3B8", lw=1.2, ls="--",
                     label="Target 1.5 m", zorder=1)
        ax_a.set_ylabel("Altitude z (m)", fontsize=10)
        ax_a.set_ylim(0, 3.6)
        ax_a.legend(loc="upper right", fontsize=8)
        ax_a.grid(True, zorder=0)

        for ax in (ax_s, ax_a):
            ax.set_xlabel("Time (s)", fontsize=10)

        status = "✓" if ep["result"] == "SUCCESS" else "✗"
        ax_s.set_title(
            f"A{ep['attempt']:02d} — {status} {ep['result']}  ({ep['duration_s']:.1f} s)",
            color=col(ep), fontsize=10.5, fontweight="bold", loc="left")

    # hide unused
    for i in range(n, nrow * ncol):
        col_i = i % ncol
        row_s = (i // ncol) * 2
        axes[row_s][col_i].set_visible(False)
        axes[row_s + 1][col_i].set_visible(False)

    plt.tight_layout(rect=[0, 0, 1, 0.978])
    plt.savefig(out)
    plt.close(fig)
    print(f"  Saved: {out.name}")


# ── Figure 5 — Obstacle proximity (safety) ────────────────────────────────────
def fig_safety(eps_done: list, meta: dict, out: Path):
    n    = len(eps_done)
    ncol = min(3, n)
    nrow = math.ceil(n / ncol)
    fig, axes = plt.subplots(nrow, ncol,
                             figsize=(ncol * 5.5, nrow * 3.8), squeeze=False)
    n_ok = sum(1 for e in meta["episodes"] if e["result"] == "SUCCESS")
    n_ep = meta["n_episodes"]
    fig.suptitle("R54 HITL — Obstacle Proximity Profile",
                 fontsize=16, fontweight="bold")
    fig.text(0.5, 0.985,
             f"RDDRONE-FMUK66 · PX4 HITL Mode  ·  {n_ok}/{n_ep} successful  ·  "
             "Red = collision zone (<0.35 m)  ·  Amber = near-miss zone (<0.80 m)",
             ha="center", fontsize=10.5, color="#475569")

    for i, ep in enumerate(eps_done):
        r, c_i = divmod(i, ncol)
        ax     = axes[r][c_i]
        df     = ep["df"]
        c      = col(ep)
        t      = df["t_s"].values
        d      = df["min_depth_m"].values

        ax.fill_between(t, d, alpha=0.18, color=c, zorder=1)
        ax.plot(t, d, color=c, lw=1.9, zorder=2, label="Min obstacle dist.")
        ax.axhline(0.80, color=C_TIMEOUT, lw=1.3, ls="--",
                   label="Near-miss (0.80 m)", zorder=3)
        ax.axhline(0.35, color=C_CRASH, lw=1.3, ls=":",
                   label="Collision (0.35 m)", zorder=3)
        ax.fill_between(t, d, 0.35,
                        where=(d < 0.35), color=C_CRASH, alpha=0.22, zorder=0)
        ax.fill_between(t, d, 0.80,
                        where=((d >= 0.35) & (d < 0.80)),
                        color=C_TIMEOUT, alpha=0.15, zorder=0)

        status = "✓ SUCCESS" if ep["result"] == "SUCCESS" else f"✗ {ep['result']}"
        ax.set_title(f"A{ep['attempt']:02d} — {status}", color=c,
                     fontsize=10.5, fontweight="bold", pad=5)
        ax.set_xlabel("Time (s)", fontsize=10)
        ax.set_ylabel("Min dist. (m)", fontsize=10)
        ax.set_ylim(bottom=-0.05)
        ax.legend(loc="upper right", fontsize=8)
        ax.grid(True, zorder=0)

    for i in range(n, nrow * ncol):
        r, c_i = divmod(i, ncol)
        axes[r][c_i].set_visible(False)

    plt.tight_layout(rect=[0, 0, 1, 0.97])
    plt.savefig(out)
    plt.close(fig)
    print(f"  Saved: {out.name}")


# ── Excel workbook ─────────────────────────────────────────────────────────────
H_FILL  = PatternFill("solid", fgColor="2563EB")
S_FILL  = PatternFill("solid", fgColor="DCFCE7")
W_FILL  = PatternFill("solid", fgColor="FEF3C7")
D_FILL  = PatternFill("solid", fgColor="FEE2E2")
Z_FILL  = PatternFill("solid", fgColor="F1F5F9")
A_FILL  = PatternFill("solid", fgColor="F8FAFF")
W_FONT  = Font(color="FFFFFF", bold=True, size=11)
B_FONT  = Font(bold=True, color="1E293B", size=10)
N_FONT  = Font(color="1E293B", size=10)
CENTER  = Alignment(horizontal="center", vertical="center", wrap_text=True)
BORDER  = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"))


def _hdr(ws, row, cols):
    for ci, v in enumerate(cols, 1):
        c = ws.cell(row, ci, v)
        c.fill = H_FILL; c.font = W_FONT
        c.alignment = CENTER; c.border = BORDER


def _outcome_fill(r: str):
    if r == "SUCCESS":    return S_FILL
    if "CRASH" in r:      return D_FILL
    return W_FILL


def build_excel(meta: dict, episodes: list, out: Path,
                speeds_avg: list, speeds_peak: list):
    wb = openpyxl.Workbook()

    # ── Sheet 1 — HITL Summary ────────────────────────────────────────────────
    ws = wb.active
    ws.title = "HITL Summary"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 32

    def _merge_title(r, val):
        ws.merge_cells(f"A{r}:B{r}")
        c = ws.cell(r, 1, val)
        c.font = Font(bold=True, size=14, color="1D4ED8")
        c.alignment = CENTER
        c.fill = PatternFill("solid", fgColor="EFF6FF")
        c.border = BORDER

    def _kv(r, k, v, vfill=None):
        ck = ws.cell(r, 1, k)
        ck.font = B_FONT; ck.alignment = CENTER
        ck.fill = A_FILL; ck.border = BORDER
        cv = ws.cell(r, 2, v)
        cv.font = N_FONT; cv.alignment = CENTER; cv.border = BORDER
        if vfill: cv.fill = vfill

    _merge_title(1, "R54 Best Policy — Hardware-in-the-Loop Test Results")
    _kv(3,  "Test Name",              meta.get("test_name", "R54 HITL"))
    _kv(4,  "Testing World",          meta.get("testing_world", ""))
    _kv(5,  "Training World",         meta.get("training_world", ""))
    _kv(6,  "Policy File",            meta.get("policy_file", ""))
    _kv(7,  "Policy Episode",         meta.get("policy_episode", ""))
    _kv(8,  "Max Speed Cap (m/s)",    meta.get("test_max_speed", 4.5))
    _kv(9,  "Corridor Half-width (m)",meta.get("corridor_hw_m", 2.5))
    _kv(10, "HITL Mode",              str(meta.get("hitl_mode", True)))
    _kv(11, "FC Firmware",            meta.get("fc_firmware", ""))
    _kv(12, "Test Timestamp",         meta.get("timestamp", ""))

    ws.cell(13, 1).fill = Z_FILL; ws.merge_cells("A13:B13")
    ws.row_dimensions[13].height = 8

    _kv(14, "Episodes Run",           meta["n_episodes"])
    _kv(15, "Successes",              meta["n_success"])
    sr = meta["success_rate"]
    _kv(16, "Success Rate",           f"{sr*100:.1f}%",
        S_FILL if sr >= 0.5 else W_FILL)
    _kv(17, "Avg Speed — CSV-derived (m/s)",
        f"{np.mean(speeds_avg):.3f}" if speeds_avg else "—")
    _kv(18, "Avg Collisions / ep",    f"{meta['avg_collisions']:.1f}")
    for r in range(3, 19):
        ws.row_dimensions[r].height = 20

    # ── Sheet 2 — Episode Log ─────────────────────────────────────────────────
    ws2 = wb.create_sheet("Episode Log")
    ws2.sheet_view.showGridLines = False
    hdrs = ["Attempt", "Outcome", "Duration (s)", "Avg Speed (m/s)",
            "Peak Speed (m/s)", "Collisions", "Near-misses",
            "Min Depth (m)", "Noise σ", "Spawn Y (m)", "Steps"]
    _hdr(ws2, 1, hdrs)
    for ci, w in enumerate([10,14,14,16,16,12,13,14,10,12,9], 1):
        ws2.column_dimensions[get_column_letter(ci)].width = w

    for ri, (ep, spd_a, spd_p) in enumerate(
            zip(meta["episodes"], speeds_avg, speeds_peak), 2):
        af = A_FILL if ri % 2 == 0 else None
        row = [ep["attempt"], ep["result"], ep["duration_s"],
               round(spd_a, 3), round(spd_p, 3),
               ep["collisions"], ep["near_misses"],
               ep["min_depth_m"], ep["dr_noise_sd"],
               ep["dr_spawn_y_m"], ep.get("steps_logged", "")]
        for ci, v in enumerate(row, 1):
            c = ws2.cell(ri, ci, v)
            c.alignment = CENTER; c.border = BORDER; c.font = N_FONT
            if ci == 2:
                c.fill = _outcome_fill(ep["result"])
                c.font = Font(bold=True, size=10, color="1E293B")
            elif af:
                c.fill = af
        ws2.row_dimensions[ri].height = 18

    # ── Sheets 3-N — Per-episode trajectory ───────────────────────────────────
    for ep in episodes:
        df = ep.get("df")
        if df is None or df.empty:
            continue
        sname = f"A{ep['attempt']:02d}_{ep['result'][:4]}"[:31]
        we = wb.create_sheet(sname)
        we.sheet_view.showGridLines = False
        _hdr(we, 1, list(df.columns))
        for ci, cname in enumerate(df.columns, 1):
            we.column_dimensions[get_column_letter(ci)].width = max(10, len(cname) + 2)
        for ri, row in enumerate(df.itertuples(index=False), 2):
            af = A_FILL if ri % 2 == 0 else None
            for ci, v in enumerate(row, 1):
                c = we.cell(ri, ci, v)
                c.alignment = Alignment(horizontal="center")
                c.border = BORDER; c.font = N_FONT
                if af: c.fill = af

    wb.save(out)
    print(f"  Saved: {out.name}")


# ── Main ───────────────────────────────────────────────────────────────────────
def main():
    p = argparse.ArgumentParser()
    p.add_argument("--summary", default=None)
    args = p.parse_args()

    summary_path = Path(args.summary) if args.summary else find_latest_summary()
    print(f"\nAnalysing: {summary_path.name}")

    with open(summary_path) as f:
        meta = json.load(f)

    eps = meta["episodes"]
    for ep in eps:
        ep["df"] = load_csv(ep.get("csv_log", ""))

    # Derive speed from CSV (more reliable than summary field)
    speeds_avg, speeds_peak = [], []
    for ep in eps:
        df = ep.get("df")
        if df is not None and not df.empty:
            s = df["speed_ms"].values
            speeds_avg.append(float(np.mean(s)))
            speeds_peak.append(float(np.max(s)))
        else:
            speeds_avg.append(ep.get("avg_speed_ms", 0.0))
            speeds_peak.append(ep.get("peak_speed_ms", 0.0))

    eps_done = [ep for ep in eps if ep.get("df") is not None]

    ts   = meta.get("timestamp", "notime")
    n_ep = meta["n_episodes"]
    stem = f"hitl_r54_{n_ep}ep_{ts}"
    out  = summary_path.parent

    print(f"\nSuccess rate : {meta['success_rate']*100:.0f}%  "
          f"({meta['n_success']}/{n_ep})")
    print(f"Avg speed    : {np.mean(speeds_avg):.3f} m/s  (CSV-derived)")

    print("\nGenerating figures...")
    fig_mission_progress(eps_done, meta, out / f"{stem}_mission_progress.png")
    fig_metrics(meta, out / f"{stem}_metrics.png")
    fig_trajectories(eps_done, meta, out / f"{stem}_trajectories.png")
    fig_speed_altitude(eps_done, meta, out / f"{stem}_speed_altitude.png")
    fig_safety(eps_done, meta, out / f"{stem}_safety.png")

    print("\nGenerating Excel workbook...")
    build_excel(meta, eps, out / f"{stem}_results.xlsx",
                speeds_avg, speeds_peak)

    print(f"\n{'='*60}")
    print(f"All outputs saved to:  {out}")
    print(f"  {stem}_mission_progress.png")
    print(f"  {stem}_metrics.png")
    print(f"  {stem}_trajectories.png")
    print(f"  {stem}_speed_altitude.png")
    print(f"  {stem}_safety.png")
    print(f"  {stem}_results.xlsx")
    print(f"{'='*60}\n")


if __name__ == "__main__":
    main()
